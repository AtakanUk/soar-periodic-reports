import base64
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from soar_periodic_reports.analytics import OffenseRow, SlaDetail, SlaSummary
from soar_periodic_reports.config import ConfigError, MailSettings
from soar_periodic_reports.emailer import build_message, xoauth2_string
from soar_periodic_reports.excel_writer import write_sla_workbook
from soar_periodic_reports.renderers import render_offense_summary
from soar_periodic_reports.soar_client import SoarClient


def test_xoauth2_string():
    decoded = base64.b64decode(xoauth2_string("bot@example.com", "tok")).decode()
    assert decoded == "user=bot@example.com\x01auth=Bearer tok\x01\x01"


@pytest.fixture
def mail_env(monkeypatch):
    for key, value in {
        "SMTP_HOST": "smtp.example.com",
        "MAIL_FROM": "bot@example.com",
        "MAIL_TO": "soc@example.com, lead@example.com",
        "WEEKLY_MAIL_TO": "weekly@example.com",
    }.items():
        monkeypatch.setenv(key, value)


def test_mail_settings_prefix_fallback(mail_env):
    assert MailSettings.from_env("WEEKLY_").recipients == ("weekly@example.com",)
    assert MailSettings.from_env("DAILY_").recipients == ("soc@example.com", "lead@example.com")
    assert MailSettings.from_env().auth == "none"


def test_oauth_requires_credentials(mail_env, monkeypatch):
    monkeypatch.setenv("SMTP_AUTH", "oauth2")
    with pytest.raises(ConfigError, match="SMTP_OAUTH_TENANT_ID"):
        MailSettings.from_env()


def test_message_with_attachment(mail_env, tmp_path):
    attachment = tmp_path / "report.xlsx"
    attachment.write_bytes(b"data")
    message = build_message(
        MailSettings.from_env(), "Subject", html_body="<p>hi</p>", attachments=[attachment]
    )
    assert message["To"] == "soc@example.com, lead@example.com"
    assert [part.get_filename() for part in message.iter_attachments()] == ["report.xlsx"]


def test_offense_summary_html_totals():
    html = render_offense_summary([OffenseRow("Brute <Force>", 3, 1, 1, 1)], hours=24)
    assert "Count - 3" in html and "No Action - 1" in html
    assert "Brute &lt;Force&gt;" in html


def test_sla_workbook(tmp_path):
    path = write_sla_workbook(
        tmp_path / "sla.xlsx",
        [SlaSummary("alice", 2, 1, timedelta(minutes=90))],
        [SlaDetail("alice", 7, "https://soar.example.com/#incidents/7", timedelta(hours=3), True)],
        threshold_minutes=120,
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Details"]
    assert [c.value for c in workbook["Summary"][2]] == ["alice", 2, 1, "01:30:00"]
    assert workbook["Details"]["C2"].hyperlink.target == "https://soar.example.com/#incidents/7"


class FakeRest:
    def __init__(self, total):
        self.items = [{"id": i} for i in range(total)]
        self.calls = []

    def post(self, uri, payload):
        self.calls.append(payload["start"])
        return {"data": self.items[payload["start"] : payload["start"] + payload["length"]]}

    def get(self, uri):
        return {"values": [{"value": 1, "label": "False-Positive"}]}


def test_soar_paging_and_labels():
    rest = FakeRest(25)
    client = SoarClient(rest, page_size=10)
    assert len(client.incidents_created_between(0, 1)) == 25
    assert rest.calls == [0, 10, 20]
    assert client.field_labels("resolution_id") == {1: "False-Positive"}
