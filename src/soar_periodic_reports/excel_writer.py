"""Excel output for the analyst SLA report."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import SlaDetail, SlaSummary, format_timedelta

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="41729F")
EXCEEDED_FILL = PatternFill("solid", fgColor="F8D7DA")


def _write_sheet(sheet, headers: Sequence[str], rows: list[list[object]]) -> None:
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(row[index - 1])) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 80)
    sheet.freeze_panes = "A2"


def write_sla_workbook(
    path: Path, summaries: list[SlaSummary], details: list[SlaDetail], threshold_minutes: int
) -> Path:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _write_sheet(
        summary_sheet,
        [
            "Analyst",
            "Total Incidents",
            f"SLA Exceeded (> {threshold_minutes} min)",
            "Average Response Time",
        ],
        [[s.analyst, s.total, s.exceeded, format_timedelta(s.average)] for s in summaries],
    )

    detail_sheet = workbook.create_sheet("Details")
    _write_sheet(
        detail_sheet,
        ["Analyst", "Incident ID", "Incident URL", "Response Time", "SLA Exceeded"],
        [
            [
                d.analyst,
                d.incident_id,
                d.incident_url,
                format_timedelta(d.response_time),
                "Yes" if d.exceeded else "No",
            ]
            for d in details
        ],
    )
    for row_index, detail in enumerate(details, start=2):
        url_cell = detail_sheet.cell(row=row_index, column=3)
        url_cell.hyperlink = detail.incident_url
        url_cell.style = "Hyperlink"
        if detail.exceeded:
            for column in range(1, 6):
                if column != 3:
                    detail_sheet.cell(row=row_index, column=column).fill = EXCEEDED_FILL

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
